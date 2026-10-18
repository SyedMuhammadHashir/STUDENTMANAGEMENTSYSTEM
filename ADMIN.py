import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import hashlib
import os
import re
import tempfile

# ============================================================
# SETTINGS
# ============================================================
ADMIN_FILE = "admin_accounts.xlsx"
DATA_ROOT = "admin_data"

STUDENT_FILE = None
WORD_FILE = None
CURRENT_ADMIN = None

BG = "#F4F7FB"
WHITE = "#FFFFFF"
DARK_BLUE = "#172B4D"
BLUE = "#1976D2"
LIGHT_BLUE = "#E3F2FD"
GREEN = "#16A085"
LIGHT_GREEN = "#E8F8F5"
RED = "#E74C3C"
LIGHT_RED = "#FDEDEC"
PURPLE = "#7E57C2"
ORANGE = "#F39C12"
DARK = "#263238"
GRAY = "#607D8B"
LIGHT_GRAY = "#ECEFF1"
CYAN = "#00BCD4"

HEADERS = ["S.No", "Roll No", "Name", "Father Name", "NIC/B-Form", "Contact Number"]
WIDTHS = {"A": 10, "B": 16, "C": 28, "D": 28, "E": 25, "F": 22}


# ============================================================
# GENERAL HELPERS
# ============================================================
def hash_password(password):
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def safe_username(username):
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", username.strip()).strip("_")
    return name or "admin"


def set_admin_workspace(username):
    """Select the separate Excel/Word files for the logged-in admin."""
    global STUDENT_FILE, WORD_FILE, CURRENT_ADMIN
    CURRENT_ADMIN = username.strip()
    safe_name = safe_username(CURRENT_ADMIN)
    unique_id = hashlib.sha256(CURRENT_ADMIN.lower().encode("utf-8")).hexdigest()[:8]
    folder = os.path.join(DATA_ROOT, f"{safe_name}_{unique_id}")
    os.makedirs(folder, exist_ok=True)
    STUDENT_FILE = os.path.join(folder, f"{safe_name}_student_records.xlsx")
    WORD_FILE = os.path.join(folder, f"{safe_name}_student_records.docx")
    create_student_file()
    if not os.path.exists(WORD_FILE):
        doc = Document()
        doc.add_heading("STUDENT MANAGEMENT SYSTEM", 0)
        doc.add_paragraph(f"Admin account: {CURRENT_ADMIN}")
        doc.add_paragraph("Use the Word File option in the application to generate the latest report.")
        doc.save(WORD_FILE)


def safe_save_workbook(workbook, path):
    """
    Save through a temporary file so a failed save does not leave a half-written
    workbook. If Excel has locked the file, show a helpful PermissionError.
    """
    folder = os.path.dirname(os.path.abspath(path)) or "."
    temp_path = None
    try:
        fd, temp_path = tempfile.mkstemp(prefix="sms_", suffix=".xlsx", dir=folder)
        os.close(fd)
        workbook.save(temp_path)
        # On Windows, os.replace fails if the target workbook is open/locked.
        os.replace(temp_path, path)
        temp_path = None
    except PermissionError as exc:
        raise PermissionError(
            f"Windows denied access to this file:\n{path}\n\n"
            "Close this admin's Excel workbook if it is open in Microsoft Excel, "
            "WPS Office, or another program. Then try again. Also check that the "
            "project folder is not read-only."
        ) from exc
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass


def style_student_sheet(ws):
    # Remove title-row merges safely, then recreate the expected title.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row <= 1 <= merged.max_row:
            ws.unmerge_cells(str(merged))
    ws.merge_cells("A1:F1")
    ws["A1"] = "STUDENT MANAGEMENT SYSTEM"
    ws["A1"].fill = PatternFill("solid", fgColor="FF172B4D")
    ws["A1"].font = Font(bold=True, size=18, color="FFFFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    thin = Side(style="thin", color="FF9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, heading in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=heading)
        cell.fill = PatternFill("solid", fgColor="FF1976D2")
        cell.font = Font(bold=True, size=12, color="FFFFFFFF")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 32

    for row in range(3, ws.max_row + 1):
        ws.row_dimensions[row].height = 30
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = PatternFill(
                "solid",
                fgColor=("FFE3F2FD" if row % 2 == 1 else "FFFFFFFF")
            )

    for column, width in WIDTHS.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:F{max(2, ws.max_row)}"


def create_student_file():
    if not STUDENT_FILE:
        return
    if not os.path.exists(STUDENT_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Student Records"
        style_student_sheet(ws)
        safe_save_workbook(wb, STUDENT_FILE)
        wb.close()
    else:
        # Only repair the file if it can be opened. Do not overwrite it here.
        try:
            wb = load_workbook(STUDENT_FILE)
            if "Student Records" not in wb.sheetnames:
                ws = wb.create_sheet("Student Records")
            else:
                ws = wb["Student Records"]
            if ws["A1"].value is None and ws.max_row <= 1:
                style_student_sheet(ws)
                safe_save_workbook(wb, STUDENT_FILE)
            wb.close()
        except PermissionError:
            raise
        except Exception as exc:
            raise RuntimeError(f"Could not open student workbook:\n{STUDENT_FILE}\n\n{exc}") from exc


def get_students():
    if not STUDENT_FILE:
        return []
    create_student_file()
    wb = load_workbook(STUDENT_FILE, data_only=True)
    try:
        ws = wb["Student Records"]
        return [list(row) for row in ws.iter_rows(min_row=3, max_col=6, values_only=True)
                if row[0] is not None]
    finally:
        wb.close()


# ============================================================
# ADMIN ACCOUNT STORAGE
# ============================================================
def create_admin_file():
    if os.path.exists(ADMIN_FILE):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Admin"
    ws.append(["Username", "Password Hash"])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="FF172B4D")
        cell.font = Font(bold=True, color="FFFFFFFF", size=12)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="FF000000"),
            right=Side(style="thin", color="FF000000"),
            top=Side(style="thin", color="FF000000"),
            bottom=Side(style="thin", color="FF000000"),
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:B1"
    safe_save_workbook(wb, ADMIN_FILE)
    wb.close()


def admin_exists():
    create_admin_file()
    wb = load_workbook(ADMIN_FILE, data_only=True)
    try:
        ws = wb["Admin"]
        return any(row[0] and str(row[0]).strip()
                   for row in ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()


def save_admin(username, password):
    username = username.strip()
    if not username or not password:
        return False, "Username and password are required."
    create_admin_file()
    wb = load_workbook(ADMIN_FILE)
    try:
        ws = wb["Admin"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and str(row[0]).strip().lower() == username.lower():
                return False, "Username already exists."
        ws.append([username, hash_password(password)])
        ws.auto_filter.ref = f"A1:B{ws.max_row}"
        safe_save_workbook(wb, ADMIN_FILE)
    except PermissionError as exc:
        return False, str(exc)
    except Exception as exc:
        return False, f"Could not save admin account:\n{exc}"
    finally:
        wb.close()

    # Create that admin's independent workspace after the account is saved.
    try:
        old_student, old_word, old_admin = STUDENT_FILE, WORD_FILE, CURRENT_ADMIN
        set_admin_workspace(username)
        # Restore the active workspace; the new admin will select it on login.
        globals()["STUDENT_FILE"] = old_student
        globals()["WORD_FILE"] = old_word
        globals()["CURRENT_ADMIN"] = old_admin
        return True, "Admin account created with separate Excel and Word files."
    except Exception as exc:
        return False, f"Account was saved, but its data files could not be created:\n{exc}"


def check_login(username, password):
    create_admin_file()
    wb = load_workbook(ADMIN_FILE, data_only=True)
    try:
        ws = wb["Admin"]
        password_hash = hash_password(password)
        for saved_username, saved_hash in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if not saved_username or not saved_hash:
                continue
            if (str(saved_username).strip().lower() == username.strip().lower()
                    and str(saved_hash) == password_hash):
                return True
        return False
    finally:
        wb.close()


# ============================================================
# WORD EXPORT
# ============================================================
def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def export_students_to_word():
    students = get_students()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("STUDENT MANAGEMENT SYSTEM")
    run.bold = True
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor(23, 43, 77)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Admin: {CURRENT_ADMIN}  |  Total Students: {len(students)}")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(25, 118, 210)

    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, heading in enumerate(HEADERS):
        cell = table.rows[0].cells[i]
        cell.text = heading
        set_cell_shading(cell, "1976D2")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(255, 255, 255)

    for index, student in enumerate(students):
        cells = table.add_row().cells
        for col, value in enumerate(student):
            cells[col].text = "" if value is None else str(value)
            set_cell_shading(cells[col], "E3F2FD" if index % 2 == 0 else "FFFFFF")
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[col].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8)

    if not students:
        cells = table.add_row().cells
        cells[0].merge(cells[-1])
        cells[0].text = "No student records available."
        set_cell_shading(cells[0], "FDEDEC")

    doc.save(WORD_FILE)
    return WORD_FILE, len(students)


# ============================================================
# TKINTER APPLICATION
# ============================================================
class StudentManagementSystem:
    
    def __init__(self, root):
        self.root = root
        self.root.title("Student Management System")
        self.root.geometry("1250x750")
        self.root.minsize(1000, 650)
        try:
            self.root.state("zoomed")
        except tk.TclError:
            pass
        self.root.configure(bg=BG)
        create_admin_file()
        self.show_login() if admin_exists() else self.show_create_admin()

    def clear_screen(self):
        for widget in self.root.winfo_children():
            widget.destroy()

    def outer(self):
        outer = tk.Frame(self.root, bg=PURPLE)
        outer.pack(fill="both", expand=True, padx=8, pady=8)
        inner = tk.Frame(outer, bg=WHITE)
        inner.pack(fill="both", expand=True, padx=4, pady=4)
        return inner

    def header(self, parent, title, color=DARK_BLUE, height=110):
        frame = tk.Frame(parent, bg=color, height=height)
        frame.pack(fill="x")
        frame.pack_propagate(False)
        tk.Label(frame, text=title, font=("Segoe UI", 24, "bold"),
                 bg=color, fg=WHITE).pack(expand=True)
        return frame

    def show_login(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "STUDENT MANAGEMENT SYSTEM", DARK_BLUE, 150)
        tk.Label(outer, text="Secure • Simple • Professional Student Records",
                 font=("Segoe UI", 13), bg=DARK_BLUE, fg="#80DEEA").place(relx=.5, y=105, anchor="center")
        center = tk.Frame(outer, bg=BG)
        center.pack(fill="both", expand=True)
        box = tk.Frame(center, bg=WHITE, highlightbackground=BLUE, highlightthickness=4)
        box.place(relx=.5, rely=.5, anchor="center", width=550, height=500)
        tk.Label(box, text="🔐 ADMIN LOGIN", font=("Segoe UI", 23, "bold"),
                 bg=WHITE, fg=DARK_BLUE).pack(pady=(35, 25))
        tk.Label(box, text="Username", font=("Segoe UI", 12, "bold"), bg=WHITE).pack(anchor="w", padx=65)
        user = tk.Entry(box, font=("Segoe UI", 14), bd=2, relief="solid")
        user.pack(fill="x", padx=65, pady=(5, 18), ipady=7)
        tk.Label(box, text="Password", font=("Segoe UI", 12, "bold"), bg=WHITE).pack(anchor="w", padx=65)
        pwd = tk.Entry(box, font=("Segoe UI", 14), show="*", bd=2, relief="solid")
        pwd.pack(fill="x", padx=65, pady=(5, 25), ipady=7)

        def login():
            username, password = user.get().strip(), pwd.get()
            if not username or not password:
                messagebox.showwarning("Required", "Enter both username and password.")
                return
            try:
                if check_login(username, password):
                    set_admin_workspace(username)
                    messagebox.showinfo("Login Successful", f"Welcome, {username}!")
                    self.show_dashboard()
                else:
                    messagebox.showerror("Login Failed", "Incorrect username or password.")
                    pwd.delete(0, tk.END)
                    pwd.focus()
            except Exception as exc:
                messagebox.showerror("File Error", str(exc))

        tk.Button(box, text="LOGIN", command=login, font=("Segoe UI", 14, "bold"),
                  bg=BLUE, fg=WHITE, relief="flat", cursor="hand2").pack(fill="x", padx=65, ipady=10)
        tk.Button(box, text="➕ CREATE ADMIN ACCOUNT", command=self.show_create_admin,
                  font=("Segoe UI", 11, "bold"), bg=LIGHT_GREEN, fg=GREEN,
                  relief="flat", cursor="hand2").pack(pady=(18, 10), ipadx=20, ipady=7)
        
        # Quit button added to login screen
        tk.Button(box, text="✖ QUIT", command=self.root.destroy, font=("Segoe UI", 11, "bold"),
                  bg=LIGHT_RED, fg=RED, relief="flat", cursor="hand2").pack(ipadx=20, ipady=6)

        user.bind("<Return>", lambda e: pwd.focus())
        pwd.bind("<Return>", lambda e: login())
        user.focus()

    def show_create_admin(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "🔐 CREATE ADMIN ACCOUNT", GREEN, 120)
        box = tk.Frame(outer, bg=WHITE, highlightbackground=GREEN, highlightthickness=4)
        box.place(relx=.5, rely=.54, anchor="center", width=600, height=520)
        entries = {}
        for label, show in [("Username", ""), ("Password", "*"), ("Confirm Password", "*")]:
            tk.Label(box, text=label, font=("Segoe UI", 12, "bold"), bg=WHITE).pack(anchor="w", padx=70, pady=(20 if not entries else 10, 0))
            ent = tk.Entry(box, font=("Segoe UI", 14), show=show, bd=2, relief="solid")
            ent.pack(fill="x", padx=70, pady=(5, 8), ipady=7)
            entries[label] = ent

        def create():
            username = entries["Username"].get().strip()
            password = entries["Password"].get()
            confirm = entries["Confirm Password"].get()
            if not username or not password:
                messagebox.showwarning("Required", "Username and password are required.")
                return
            if len(password) < 4:
                messagebox.showwarning("Password", "Password must contain at least 4 characters.")
                return
            if password != confirm:
                messagebox.showerror("Password Error", "Passwords do not match.")
                return
            success, msg = save_admin(username, password)
            if success:
                messagebox.showinfo("Account Created", msg + "\n\nYou can now log in.")
                self.show_login()
            else:
                messagebox.showerror("Account Error", msg)

        tk.Button(box, text="CREATE ADMIN ACCOUNT", command=create, font=("Segoe UI", 13, "bold"),
                  bg=GREEN, fg=WHITE, relief="flat", cursor="hand2").pack(fill="x", padx=70, pady=18, ipady=10)
        tk.Button(box, text="← BACK TO LOGIN", command=self.show_login, font=("Segoe UI", 11, "bold"),
                  bg=LIGHT_GRAY, fg=DARK, relief="flat", cursor="hand2").pack(ipadx=20, ipady=6)
        
        # Quit button
        tk.Button(box, text="✖ QUIT", command=self.root.destroy, font=("Segoe UI", 11, "bold"),
                  bg=LIGHT_RED, fg=RED, relief="flat", cursor="hand2").pack(pady=(10, 0), ipadx=20, ipady=6)
        
        entries["Username"].focus()

    def show_dashboard(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, f"🎓 STUDENT MANAGEMENT SYSTEM — {CURRENT_ADMIN}", DARK_BLUE, 125)
        tk.Label(outer, text="ADMIN CONTROL PANEL", font=("Segoe UI", 11, "bold"),
                 bg=DARK_BLUE, fg="#80DEEA").place(relx=.5, y=91, anchor="center")
        content = tk.Frame(outer, bg=BG)
        content.pack(fill="both", expand=True, padx=35, pady=25)
        actions = [
            ("👨‍🎓", "STUDENT RECORD", "Add New Student", BLUE, self.show_student_record),
            ("📋", "SEE RECORD", "View All Students", GREEN, self.show_records),
            ("🔎", "SEARCH RECORD", "Find Student", PURPLE, self.show_search),
            ("📄", "WORD FILE", "Create Colorful Report", CYAN, self.export_word_file),
            ("🗑️", "DELETE RECORD", "Remove Student", RED, self.show_delete),
            ("🚪", "LOG OUT", "Return to Login", ORANGE, self.logout),
        ]
        for i, (emoji, title, sub, color, command) in enumerate(actions):
            r, c = divmod(i, 3)
            card = tk.Frame(content, bg=WHITE, highlightbackground=color, highlightthickness=3)
            card.grid(row=r, column=c, padx=15, pady=15, sticky="nsew")
            content.grid_rowconfigure(r, weight=1)
            content.grid_columnconfigure(c, weight=1)
            tk.Label(card, text=emoji, font=("Segoe UI Emoji", 32), bg=WHITE).pack(pady=(15, 3))
            tk.Label(card, text=title, font=("Segoe UI", 14, "bold"), bg=WHITE, fg=color).pack()
            tk.Label(card, text=sub, font=("Segoe UI", 10), bg=WHITE, fg=GRAY).pack(pady=4)
            tk.Button(card, text="OPEN", command=command, bg=color, fg=WHITE,
                      font=("Segoe UI", 10, "bold"), relief="flat", cursor="hand2").pack(
                          padx=40, pady=12, fill="x", ipady=6)
        tk.Label(outer, text=f"",
                 font=("Segoe UI", 9), bg=WHITE, fg=GRAY).pack(pady=7)

    def logout(self):
        global STUDENT_FILE, WORD_FILE, CURRENT_ADMIN
        STUDENT_FILE = WORD_FILE = CURRENT_ADMIN = None
        self.show_login()

    def show_student_record(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "👨‍🎓 STUDENT RECORD", BLUE, 95)
        form = tk.Frame(outer, bg=BG, highlightbackground=BLUE, highlightthickness=3)
        form.pack(fill="both", expand=True, padx=50, pady=30)
        fields = ["Roll No", "Name", "Father Name", "NIC/B-Form", "Contact Number"]
        entries = {}
        for i, field in enumerate(fields):
            r, c = divmod(i, 2)
            area = tk.Frame(form, bg=BG)
            area.grid(row=r, column=c, padx=35, pady=14, sticky="ew")
            form.grid_columnconfigure(c, weight=1)
            tk.Label(area, text=field, font=("Segoe UI", 12, "bold"), bg=BG).pack(anchor="w")
            entry = tk.Entry(area, font=("Segoe UI", 13), bd=2, relief="solid")
            entry.pack(fill="x", pady=5, ipady=7)
            entries[field] = entry

        def add_student():
            values = {key: ent.get().strip() for key, ent in entries.items()}
            for key, value in values.items():
                if not value:
                    messagebox.showwarning("Required", f"Please enter {key}.")
                    entries[key].focus()
                    return
            try:
                wb = load_workbook(STUDENT_FILE)
                try:
                    ws = wb["Student Records"]
                    for row in ws.iter_rows(min_row=3, values_only=True):
                        if row[1] is not None and str(row[1]).strip() == values["Roll No"]:
                            messagebox.showerror("Duplicate Roll No", "This Roll No already exists.")
                            return
                    serial = max(1, ws.max_row - 1)
                    ws.append([serial, values["Roll No"], values["Name"], values["Father Name"],
                               values["NIC/B-Form"], values["Contact Number"]])
                    style_student_sheet(ws)
                    safe_save_workbook(wb, STUDENT_FILE)
                finally:
                    wb.close()
                messagebox.showinfo("Success", "Student record added successfully.")
                for ent in entries.values():
                    ent.delete(0, tk.END)
                entries["Roll No"].focus()
            except Exception as exc:
                messagebox.showerror("Excel Error", f"Could not save record:\n\n{exc}")

        tk.Button(form, text="➕ ADD STUDENT RECORD", command=add_student, font=("Segoe UI", 13, "bold"),
                  bg=GREEN, fg=WHITE, relief="flat", cursor="hand2").grid(
                      row=3, column=0, padx=40, pady=25, sticky="ew", ipady=8)
        tk.Button(form, text="← BACK TO MENU", command=self.show_dashboard, font=("Segoe UI", 13, "bold"),
                  bg=DARK_BLUE, fg=WHITE, relief="flat", cursor="hand2").grid(
                      row=3, column=1, padx=40, pady=25, sticky="ew", ipady=8)
        entries["Roll No"].focus()

    def configure_treeview(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Treeview", background=WHITE, foreground=DARK, rowheight=32,
                        fieldbackground=WHITE, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", background=BLUE, foreground=WHITE,
                        font=("Segoe UI", 10, "bold"), padding=7)
        style.map("Treeview", background=[("selected", PURPLE)], foreground=[("selected", WHITE)])

    def make_tree(self, parent):
        columns = tuple(HEADERS)
        tree = ttk.Treeview(parent, columns=columns, show="headings")
        widths = [70, 120, 200, 200, 180, 170]
        for column, width in zip(columns, widths):
            tree.heading(column, text=column)
            tree.column(column, width=width, anchor="center")
        return tree

    def show_records(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "📋 ALL STUDENT RECORDS", DARK_BLUE, 75)
        frame = tk.Frame(outer, bg=BG)
        frame.pack(fill="both", expand=True, padx=25, pady=20)
        self.configure_treeview()
        tree = self.make_tree(frame)
        vs = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        hs = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        try:
            students = get_students()
            for student in students:
                tree.insert("", "end", values=student)
            tk.Label(outer, text=f"Total Students: {len(students)}", font=("Segoe UI", 11, "bold"),
                     bg=WHITE, fg=GREEN).pack(pady=4)
        except Exception as exc:
            messagebox.showerror("Excel Error", str(exc))
        tk.Button(outer, text="← BACK TO MENU", command=self.show_dashboard, bg=DARK_BLUE,
                  fg=WHITE, font=("Segoe UI", 12, "bold"), relief="flat").pack(pady=8, ipadx=25, ipady=6)

    def show_search(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "🔎 SEARCH STUDENT RECORD", PURPLE, 75)
        search_area = tk.Frame(outer, bg=BG)
        search_area.pack(fill="x", padx=35, pady=15)
        tk.Label(search_area, text="Search by roll number, name, father name, NIC/B-Form or contact",
                 font=("Segoe UI", 11, "bold"), bg=BG).pack(pady=5)
        entry = tk.Entry(search_area, font=("Segoe UI", 13), bd=2, relief="solid")
        entry.pack(fill="x", padx=120, ipady=7)
        self.configure_treeview()
        tree = self.make_tree(outer)
        tree.pack(fill="both", expand=True, padx=30, pady=10)

        def search():
            for item in tree.get_children():
                tree.delete(item)
            keyword = entry.get().strip().lower()
            if not keyword:
                messagebox.showwarning("Search", "Enter something to search.")
                return
            try:
                students = get_students()
                matches = [s for s in students if any(keyword in str(v).lower() for v in s if v is not None)]
                for student in matches:
                    tree.insert("", "end", values=student)
                if not matches:
                    messagebox.showinfo("Search Result", "No student record found.")
            except Exception as exc:
                messagebox.showerror("Excel Error", str(exc))

        tk.Button(search_area, text="SEARCH", command=search, bg=PURPLE, fg=WHITE,
                  font=("Segoe UI", 11, "bold"), relief="flat").pack(pady=8, ipadx=25, ipady=5)
        entry.bind("<Return>", lambda e: search())
        tk.Button(outer, text="← BACK TO MENU", command=self.show_dashboard, bg=DARK_BLUE,
                  fg=WHITE, font=("Segoe UI", 12, "bold"), relief="flat").pack(pady=8, ipadx=25, ipady=6)
        entry.focus()

    def show_delete(self):
        self.clear_screen()
        outer = self.outer()
        self.header(outer, "🗑️ DELETE STUDENT RECORD", RED, 75)
        box = tk.Frame(outer, bg=WHITE, highlightbackground=RED, highlightthickness=4)
        box.place(relx=.5, rely=.51, anchor="center", width=600, height=360)
        tk.Label(box, text="Enter Student Roll Number", font=("Segoe UI", 16, "bold"),
                 bg=WHITE, fg=DARK).pack(pady=(45, 15))
        roll_entry = tk.Entry(box, font=("Segoe UI", 15), bd=2, relief="solid")
        roll_entry.pack(padx=70, fill="x", ipady=9)

        def delete():
            roll = roll_entry.get().strip()
            if not roll:
                messagebox.showwarning("Required", "Please enter Roll No.")
                return
            if not messagebox.askyesno("Confirm Delete", f"Are you sure you want to delete Roll No {roll}?"):
                return
            wb = None
            try:
                wb = load_workbook(STUDENT_FILE)
                ws = wb["Student Records"]
                found_row = None
                for row in range(3, ws.max_row + 1):
                    value = ws.cell(row=row, column=2).value
                    if value is not None and str(value).strip() == roll:
                        found_row = row
                        break
                if found_row is None:
                    messagebox.showerror("Not Found", f"No student found with Roll No {roll}.")
                    return
                ws.delete_rows(found_row, 1)
                # Re-number serial numbers and style the same workbook in memory.
                for row in range(3, ws.max_row + 1):
                    ws.cell(row=row, column=1).value = row - 2
                style_student_sheet(ws)
                # Save exactly once. Do not call a second styling/save function afterward.
                safe_save_workbook(wb, STUDENT_FILE)
                messagebox.showinfo("Deleted", "Student record deleted successfully.")
                roll_entry.delete(0, tk.END)
                roll_entry.focus()
            except Exception as exc:
                messagebox.showerror("Excel Error", f"Could not delete record:\n\n{exc}")
            finally:
                if wb is not None:
                    wb.close()

        tk.Button(box, text="🗑️ DELETE RECORD", command=delete, font=("Segoe UI", 13, "bold"),
                  bg=RED, fg=WHITE, relief="flat", cursor="hand2").pack(
                      fill="x", padx=70, pady=25, ipady=9)
        roll_entry.bind("<Return>", lambda e: delete())
        tk.Button(outer, text="← BACK TO MENU", command=self.show_dashboard, bg=DARK_BLUE,
                  fg=WHITE, font=("Segoe UI", 12, "bold"), relief="flat").pack(side="bottom", pady=15, ipadx=25, ipady=6)
        roll_entry.focus()

    def export_word_file(self):
        try:
            filename, count = export_students_to_word()
            messagebox.showinfo("Word File Created",
                                f"Word report created successfully!\n\nStudents included: {count}\nFile:\n{filename}")
        except PermissionError:
            messagebox.showerror("Word File Is Open",
                                 "Close this admin's Word report in Microsoft Word/WPS Office, then try again.")
        except Exception as exc:
            messagebox.showerror("Word Export Error", f"Could not create Word file:\n\n{exc}")


if __name__ == "__main__":
    root = tk.Tk()
    try:
        app = StudentManagementSystem(root)
        root.mainloop()
    except Exception as error:
        messagebox.showerror("Startup Error", str(error))
        root.destroy()